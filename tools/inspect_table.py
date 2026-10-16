"""看一张表里各列是怎么填的（只读）。

用法：python tools\\inspect_table.py [表路径] [行号]
不给参数就看配置里的当前表。用来核对"老表某个商品当年是用哪几列上的"。
"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

if len(sys.argv) > 1:
    table = Path(sys.argv[1])
else:
    from temu_bot.config import load_config

    table = load_config().project.table

which = int(sys.argv[2]) if len(sys.argv) > 2 else None

wb = load_workbook(table, data_only=True)
ws = wb.active
print(f"表：{table}")
print(f"工作表：{ws.title}｜最大行 {ws.max_row}｜最大列 {ws.max_column}")

header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
print("\n表头：")
for i, name in enumerate(header, start=1):
    print(f"  {i:>2}. {name}")

# 每一列有多少行是填了值的
print("\n各列填写情况：")
for c in range(1, ws.max_column + 1):
    filled = sum(
        1
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=c).value not in (None, "")
    )
    if filled:
        print(f"  第 {c:>2} 列 {str(header[c - 1])[:24]:<26} 填了 {filled} 行")

# 每一行有哪几列是填了的（看模式）
patterns = Counter()
for r in range(2, ws.max_row + 1):
    cols = tuple(
        c
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=r, column=c).value not in (None, "")
    )
    patterns[cols] += 1
print("\n行的填写模式（哪几列有值 → 有几行）：")
for cols, n in patterns.most_common(10):
    print(f"  {n:>4} 行  列 {list(cols)}")

if which:
    print(f"\n第 {which} 行逐列：")
    for c in range(1, ws.max_column + 1):
        value = ws.cell(row=which, column=c).value
        if value not in (None, ""):
            print(f"  {c:>2}. {str(header[c - 1])[:24]:<26} = {value!r}")
