"""直接读某个 xlsx 的指定格子（绕过任何缓存），看第 13/14 列到底有没有值。"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"C:\Users\A\Desktop\yyy\商品信息表.xlsx")
wb = load_workbook(path)
ws = wb.active

print(f"{path}｜工作表 {ws.title}｜{ws.max_row} 行 × {ws.max_column} 列")
print(f"{'行':>4}  {'第2列 主图':<22} {'第13列 商品ID':<22} {'第14列 状态':<8}")
for r in range(1, ws.max_row + 1):
    print(
        f"{r:>4}  {str(ws.cell(row=r, column=2).value or ''):<22}"
        f" {str(ws.cell(row=r, column=13).value or ''):<22}"
        f" {str(ws.cell(row=r, column=14).value or ''):<8}"
    )
wb.close()
