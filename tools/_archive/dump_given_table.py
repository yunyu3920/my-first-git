"""把用户给的那份表格结构原样 dump 出来（只读），用来对齐"收集图片"的产出格式。"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

import openpyxl

path = Path(sys.argv[1])
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print(f"文件: {path.name}")
print(f"工作表: {wb.sheetnames}")

for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(min_row=1, max_row=8, values_only=True))
    if not rows:
        print(f"\n=== [{name}] 空 ===")
        continue
    header = rows[0]
    print(f"\n=== [{name}] 列数={len(header)} ===")
    for i, h in enumerate(header, 1):
        print(f"  {i:2d}. {h!r}")

    print(f"\n--- [{name}] 前 {len(rows) - 1} 行数据 ---")
    for n, row in enumerate(rows[1:], start=2):
        print(f"\n  第{n}行:")
        for i, h in enumerate(header, 1):
            v = row[i - 1] if len(row) >= i else None
            if v is None or str(v).strip() == "":
                continue
            print(f"    {i:2d}. {h!r} = {str(v)!r}")

    # 全表统计
    print(f"\n--- [{name}] 全表统计 ---")
    all_rows = list(ws.iter_rows(min_row=2, values_only=True))
    print(f"  数据行数: {len(all_rows)}")
    for idx, h in enumerate(header, 1):
        vals = [r[idx - 1] for r in all_rows if len(r) >= idx and r[idx - 1] not in (None, "")]
        uniq = Counter(str(v) for v in vals)
        sample = list(uniq)[:3]
        print(f"  {idx:2d}. {h!r}: 非空 {len(vals)}/{len(all_rows)}　样例={sample}")

wb.close()
