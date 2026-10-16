"""把 xlsx 里**每个工作表**的第 13 列都打出来，看是不是有多张表、配置指错了。"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from temu_bot.config import load_config

path = Path(sys.argv[1]) if len(sys.argv) > 1 else load_config().project.table
cfg = load_config()

wb = load_workbook(path)
print(f"文件：{path}")
print(f"配置 sheet_name = {cfg.project.sheet_name!r}")
print(f"文件里的工作表：{wb.sheetnames}")
print(f"wb.active = {wb.active.title!r}\n")

for name in wb.sheetnames:
    ws = wb[name]
    ids = [
        ws.cell(row=r, column=13).value
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=13).value not in (None, "")
    ]
    print(f"工作表 {name!r}：{ws.max_row} 行，第 13 列有值的有 {len(ids)} 行")
    for v in ids[:8]:
        print(f"    {v}")
wb.close()
