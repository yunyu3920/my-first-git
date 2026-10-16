"""打印某一行的全部 20 列原值 + 标题规范化结果（只读）。"""

from __future__ import annotations

import sys

import openpyxl

from temu_bot.product import bad_chars, normalize_en_title

PATH = r"C:\Users\A\Desktop\Temu\商品信息表.xlsx"
TARGET = int(sys.argv[1]) if len(sys.argv) > 1 else 32

wb = openpyxl.load_workbook(PATH, read_only=True, data_only=True)
ws = wb["Sheet1"]
header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))

for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if idx != TARGET:
        continue
    print(f"=== 第 {idx} 行 ===")
    for i, name in enumerate(header, 1):
        val = row[i - 1] if len(row) >= i else None
        val = "" if val is None else str(val)
        mark = "  <<<" if name in ("英文标题", "中文标题", "状态", "商品ID", "问题项") else ""
        print(f"  {i:2d}. {name}: {val!r}{mark}")

    cn = str(row[16] or "")
    en = str(row[17] or "")
    print("\n--- 英文标题规范化 ---")
    print(f"  原文  : {en!r}")
    norm = normalize_en_title(en)
    print(f"  规范化: {norm!r}")
    print(f"  长度  : {len(norm)}（Temu 上限按 120 计）")
    bad = bad_chars(norm)
    print(f"  非法字符: {bad or '无'}")
    break
else:
    print(f"没有第 {TARGET} 行")

wb.close()
