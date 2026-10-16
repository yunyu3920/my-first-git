"""搞清楚"SKU 货号"到底是怎么生成的（只读）。

做两件事：
  1. 打出一个商品**所有** SKU 的 extCode（看是不是 前缀+序号/S尺码 拼出来的）；
  2. 把老表里"上架成功"的行和它商品上的 extCode 对起来，看当年那个工具是怎么填的。
"""

from __future__ import annotations

import json
import sys
from collections import Counter

from openpyxl import load_workbook

from temu_bot.config import load_config
from temu_bot.dianxiaomi.client import DxmClient

OLD_TABLE = r"C:\Users\A\Desktop\Temu\商品信息表.xlsx"


def codes_of(client: DxmClient, product_id: str) -> dict:
    product = ((client.product_edit(product_id).get("data") or {}).get("product")) or {}
    if not product:
        return {}

    skc = json.loads(product.get("mainProductSkuSpecReqs") or "[]")
    var = json.loads(product.get("variationListStr") or "[]")

    return {
        "name": product.get("productName"),
        "skc_ext": [(e.get("specName"), e.get("extCode")) for e in skc],
        "sku_ext": [
            (
                "|".join(
                    s.get("specName", "")
                    for s in json.loads(e.get("productSkuSpecReqs") or "[]")
                ),
                e.get("extCode"),
                e.get("code"),
                e.get("codeType"),
            )
            for e in var
        ],
    }


cfg = load_config()
login = sys.argv[1] if len(sys.argv) > 1 else "hat-tshirt"
product_id = sys.argv[2] if len(sys.argv) > 2 else ""

with DxmClient(cfg, login) as client:
    if product_id:
        info = codes_of(client, product_id)
        print(f"=== 商品 {product_id}｜{info.get('name')} ===")
        print(f"  SKC extCode：{info['skc_ext']}")
        print("  SKU (颜色|尺码, extCode, code, codeType)：")
        for row in info["sku_ext"]:
            print(f"    {row}")

    # 老表：前缀编码 ↔ 商品ID ↔ 商品上的 extCode
    print(f"\n=== 老表 {OLD_TABLE} ===")
    wb = load_workbook(OLD_TABLE, data_only=True)
    ws = wb.active
    print(f"{'行':>4} {'第9列 前缀编码':<16} {'第19列 SKC货号':<18} {'第13列 商品ID':<20}")
    pairs: list[tuple[int, str, str]] = []
    for r in range(2, min(ws.max_row, 12) + 1):
        c9 = ws.cell(row=r, column=9).value
        c19 = ws.cell(row=r, column=19).value
        c13 = ws.cell(row=r, column=13).value
        print(f"{r:>4} {str(c9 or ''):<16} {str(c19 or ''):<18} {str(c13 or ''):<20}")
        if c13:
            pairs.append((r, str(c9 or ""), str(c13)))

    print("\n=== 老商品上的 extCode ===")
    for r, prefix, pid in pairs[:4]:
        try:
            info = codes_of(client, pid)
        except Exception as exc:  # noqa: BLE001
            print(f"  第 {r} 行 {pid}: 读失败 {exc}")
            continue
        exts = [x[1] for x in info.get("sku_ext", [])]
        skc = [x[1] for x in info.get("skc_ext", [])]
        print(f"  第 {r} 行  表里前缀={prefix}")
        print(f"      SKC extCode = {set(skc)}")
        print(f"      SKU extCode = {Counter(exts)}")
