"""拿**上架成功过的老商品**对照：表里第 9 列的值，到底被填进了哪个字段（只读）。

老表是 A lovely costume 的，所以必须用 `lovely` 这个登录身份去读。
"""

from __future__ import annotations

import json

from openpyxl import load_workbook

from temu_bot.config import load_config
from temu_bot.dianxiaomi.client import DxmClient

OLD_TABLE = r"C:\Users\A\Desktop\Temu\商品信息表.xlsx"
LOGIN = "lovely"

cfg = load_config()

wb = load_workbook(OLD_TABLE, data_only=True)
ws = wb.active

pairs: list[tuple[int, str, str, str]] = []
for r in range(2, ws.max_row + 1):
    c9 = str(ws.cell(row=r, column=9).value or "")
    c13 = str(ws.cell(row=r, column=13).value or "")
    c14 = str(ws.cell(row=r, column=14).value or "")
    c19 = str(ws.cell(row=r, column=19).value or "")
    if c13 and c9:
        pairs.append((r, c9, c13, c19))

print(f"老表里填了第9列+商品ID 的有 {len(pairs)} 行\n")

shown = 0
with DxmClient(cfg, LOGIN) as client:
    for r, c9, c13, c19 in pairs:
        try:
            product = ((client.product_edit(c13).get("data") or {}).get("product")) or {}
        except Exception as exc:  # noqa: BLE001
            print(f"第 {r} 行 {c13}: 读失败 {exc}")
            continue
        if not product:
            print(f"第 {r} 行 {c13}: 读不到商品")
            continue

        skc = json.loads(product.get("mainProductSkuSpecReqs") or "[]")
        var = json.loads(product.get("variationListStr") or "[]")

        skc_ext = sorted({str(e.get("extCode")) for e in skc})
        sku_ext = sorted({str(e.get("extCode")) for e in var})

        print(f"第 {r} 行｜{product.get('productName')}")
        print(f"   表里第9列(前缀编码) = {c9!r}    第19列(SKC货号) = {c19!r}")
        print(f"   商品 SKC extCode   = {skc_ext}")
        print(f"   商品 SKU extCode   = {sku_ext}")
        print(f"   → 命中第9列？ SKC {'✔' if c9 in skc_ext else '✘'} / SKU {'✔' if c9 in sku_ext else '✘'}")
        print()

        shown += 1
        if shown >= 5:
            break
