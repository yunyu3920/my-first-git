"""表格占用保护与安全回写。

三条规矩（对齐用户《图片批处理工具箱》的安全约定）：
1. 写前必须确认表没被 Excel/WPS 打开（~$ 锁文件 + 独占打开试探）。
2. 每次回写前先备份（带时间戳，保留最近 N 份）。
3. 落盘走「同目录临时文件 + os.replace」原子替换，绝不半截文件覆盖原表。
"""

from __future__ import annotations

import logging
import os
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from filelock import FileLock, Timeout

from temu_bot.config import Config

log = logging.getLogger(__name__)


class TableBusyError(RuntimeError):
    """表格被占用（多半是开着 Excel/WPS），此时绝不能写。"""


@dataclass(slots=True)
class CellUpdate:
    """一行的回写内容：{列号: 值}。"""

    row_index: int
    fields: dict[int, Any] = field(default_factory=dict)

    def is_empty(self) -> bool:
        return not self.fields


# ------------------------------------------------------------------ 忙检测
def excel_lock_files(table: Path) -> list[Path]:
    """Excel/WPS 打开工作簿时会在同目录留下 ~$ 开头的锁文件。"""
    try:
        return sorted(p for p in table.parent.glob("~$*") if p.is_file())
    except OSError:
        return []


def assert_not_busy(table: Path) -> None:
    locks = excel_lock_files(table)
    if locks:
        raise TableBusyError(
            f"检测到 Office 锁文件「{locks[0].name}」——表正被 Excel/WPS 打开，"
            f"请先关闭后重试：{table}"
        )
    if not table.exists():
        raise FileNotFoundError(f"找不到表格：{table}")
    try:
        with table.open("r+b"):
            pass
    except PermissionError as exc:
        raise TableBusyError(f"表格被占用，无法独占打开（可能正在被写入）：{table}") from exc
    except OSError as exc:
        raise TableBusyError(f"表格无法打开：{table}（{exc}）") from exc


# ------------------------------------------------------------------ 备份
def backup(table: Path, backup_dir: Path, keep: int = 20) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    # 带毫秒：一秒内连写多行（批量上架）时，只精确到秒会**互相覆盖**，
    # 那样就丢了"这一批跑之前"的那份状态——正是最该留住的那份。
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")[:-3]
    dest = backup_dir / f"{table.stem}.{stamp}.bak{table.suffix}"
    shutil.copy2(table, dest)
    log.info("已备份表格 → %s", dest)

    old = sorted(backup_dir.glob(f"{table.stem}.*.bak{table.suffix}"), key=lambda p: p.stat().st_mtime)
    for stale in old[:-keep] if keep > 0 else []:
        try:
            stale.unlink()
            log.debug("清理旧备份 %s", stale.name)
        except OSError:
            log.warning("旧备份删不掉：%s", stale)
    return dest


# ------------------------------------------------------------------ 回写
def write_updates(
    cfg: Config,
    updates: list[CellUpdate],
    *,
    sheet_name: str | None = None,
    backup_first: bool = True,
) -> Path | None:
    """把若干行的若干列写回商品表，返回备份路径（没写任何东西时返回 None）。

    用 filelock 保证同一时刻只有一个进程在写表；openpyxl 正向加载并保存以保留
    单元格样式（图表/图片/数据透视这类对象 openpyxl 不保证保留，所以必须先备份）。
    """
    pending = [u for u in updates if not u.is_empty()]
    if not pending:
        log.info("没有需要回写的内容，跳过")
        return None

    table = cfg.project.table
    sheet = sheet_name or cfg.project.sheet_name
    lock_file = cfg.data_dir / "state" / "table.lock"
    lock_file.parent.mkdir(parents=True, exist_ok=True)

    try:
        with FileLock(lock_file, timeout=30):
            assert_not_busy(table)

            backup_path = None
            if backup_first:
                backup_path = backup(table, cfg.data_dir / "backups", cfg.sheet.backup_keep)

            wb = openpyxl.load_workbook(table)
            try:
                if sheet not in wb.sheetnames:
                    raise KeyError(f"表里没有工作表「{sheet}」，实际有：{wb.sheetnames}")
                ws = wb[sheet]

                for update in pending:
                    if update.row_index <= cfg.sheet.header_row:
                        raise ValueError(f"拒绝写表头行或以上：row_index={update.row_index}")
                    for col, value in update.fields.items():
                        # **不能写成 `ws.cell(row=..., column=..., value=value)`**：
                        # openpyxl 的 Worksheet.cell() 在 `value is None` 时直接跳过赋值，
                        # 于是"把格子清空"（value=None）会静默变成什么都没做——文件照样
                        # 被重存（体积还变了），看着像成功，其实值还在。踩过。
                        # 所以拿到 cell 再赋 `.value`。
                        ws.cell(row=update.row_index, column=int(col)).value = value

                stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
                tmp = table.with_name(f"{table.stem}.{stamp}.tmp{table.suffix}")
                try:
                    wb.save(tmp)
                    os.replace(tmp, table)
                finally:
                    if tmp.exists():
                        try:
                            tmp.unlink()
                        except OSError:
                            pass
            finally:
                wb.close()
    except Timeout as exc:
        raise TableBusyError(f"另一个 temu-bot 进程正在写表，等锁超时：{lock_file}") from exc

    log.info("已回写 %d 行 → %s", len(pending), table)
    return backup_path
