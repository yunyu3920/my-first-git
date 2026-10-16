"""商品信息表读取：表头校验 + 逐行产出 SheetRow。

只读路径统一走 openpyxl 的 read_only 模式（这张表有上万行也不吃内存）；
回写不走这里，见 guard.write_updates()。
"""

from __future__ import annotations

import logging
from collections import Counter
from pathlib import Path
from typing import Iterator

import openpyxl

from temu_bot.config import Config, SheetCfg
from temu_bot.sheet import schema
from temu_bot.sheet.schema import SheetRow, build_row, normalize_cell

log = logging.getLogger(__name__)


class SheetTable:
    def __init__(self, path: Path, sheet_name: str = "Sheet1", cfg: SheetCfg | None = None) -> None:
        self.path = Path(path)
        self.sheet_name = sheet_name
        self.cfg = cfg or SheetCfg()
        self._headers: dict[int, str] | None = None

    # ------------------------------------------------------------ 构造
    @classmethod
    def from_config(cls, cfg: Config) -> "SheetTable":
        return cls(cfg.project.table, cfg.project.sheet_name, cfg.sheet)

    # ------------------------------------------------------------ 原始读
    def _iter_raw(self) -> Iterator[tuple[int, dict[int, str]]]:
        """逐行产出 (Excel 行号, {列号: 文本})。行号从 1 开始。"""
        if not self.path.exists():
            raise FileNotFoundError(f"找不到商品表：{self.path}")

        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        try:
            if self.sheet_name not in wb.sheetnames:
                raise KeyError(
                    f"表里没有工作表「{self.sheet_name}」，实际有：{wb.sheetnames}"
                )
            ws = wb[self.sheet_name]
            for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                cells: dict[int, str] = {}
                for i, value in enumerate(row, start=1):
                    text = normalize_cell(value)
                    if text:
                        cells[i] = text
                yield idx, cells
        finally:
            wb.close()

    # ------------------------------------------------------------ 表头
    def read_headers(self) -> dict[int, str]:
        if self._headers is not None:
            return self._headers
        header_row = self.cfg.header_row
        for idx, cells in self._iter_raw():
            if idx == header_row:
                self._headers = cells
                break
        else:
            raise ValueError(f"表里连第 {header_row} 行表头都没有：{self.path}")
        return self._headers

    def validate_headers(self) -> list[str]:
        """表头与契约的差异（空列表 = 完全一致）。列漂移会让回写打错列，必须先查。"""
        return schema.header_mismatches(self.read_headers())

    def assert_headers_ok(self) -> None:
        problems = self.validate_headers()
        if problems:
            raise ValueError("表头与列契约不一致，拒绝继续：\n  - " + "\n  - ".join(problems))

    # ------------------------------------------------------------ 数据行
    def iter_rows(
        self,
        *,
        limit: int | None = None,
        row_key: str | None = None,
        row_index: int | None = None,
        only_status: str | None = None,
        only_shop: str | None = None,
        skip_status: tuple[str, ...] = (),
        skip_done: bool = False,
    ) -> Iterator[SheetRow]:
        """按筛选条件产出数据行。

        skip_done=True 时跳过状态为「成功」的行（默认不跳，交给调用方决定）。
        """
        header_row = self.cfg.header_row
        success = self.cfg.statuses.get("success", "成功")
        emitted = 0

        for idx, cells in self._iter_raw():
            if idx <= header_row:
                continue

            row = build_row(
                idx,
                cells,
                separator=self.cfg.multi_value_separator,
                fallback_pattern=self.cfg.separator_fallback_pattern,
            )
            if row is None:
                continue

            if row_key and row.row_key != row_key:
                continue
            if row_index and row.row_index != row_index:
                continue
            if only_status and row.status != only_status:
                continue
            if only_shop and row.shop != only_shop:
                continue
            if skip_status and row.status in skip_status:
                continue
            if skip_done and row.status == success:
                continue

            yield row
            emitted += 1
            if limit is not None and emitted >= limit:
                return

    def count(self, *, skip_done: bool = False) -> int:
        return sum(1 for _ in self.iter_rows(skip_done=skip_done))

    def stats(self) -> dict[str, object]:
        """行数 + 状态分布 + 店铺分布，用于 doctor / 控制台概览。"""
        status: Counter[str] = Counter()
        shops: Counter[str] = Counter()
        total = 0
        for row in self.iter_rows():
            total += 1
            status[row.status or "(空)"] += 1
            shops[row.shop or "(空)"] += 1
        return {
            "rows": total,
            "by_status": dict(status.most_common()),
            "by_shop": dict(shops.most_common()),
        }

    def by_row_key(self, row_key: str) -> SheetRow | None:
        for row in self.iter_rows(row_key=row_key):
            return row
        return None
