"""回写表格（`sheet/writeback.py` + 流程里的 WRITEBACK 步）。

这一步是**防重复铺货的闸门**：表格里没有商品ID 的行，下次跑还是"新建"。
所以重点测三件事：
1. 该写什么（商品ID / 状态 / 问题项）；
2. 干跑绝不写；
3. 写不成（表被占用）时**不能把上架结果带崩**，也不能把状态写成"失败"。
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from temu_bot.listing import ListingPipeline, Status, StepResult
from temu_bot.sheet.schema import (
    COL_PROBLEM,
    COL_PRODUCT_ID,
    COL_STATUS,
    SheetRow,
    make_row_key,
)
from temu_bot.sheet.table import SheetTable
from temu_bot.sheet.writeback import (
    PROBLEM_MAX,
    first_problem,
    updates_for,
    write_back,
)

STATUSES = {"success": "成功", "failed": "失败", "need_human": "需人工"}


def make_row(tmp_path: Path, index: int = 2, product_id: str = "") -> SheetRow:
    folder = tmp_path / "imgs"
    folder.mkdir(exist_ok=True)
    return SheetRow(
        row_index=index,
        row_key=make_row_key(str(folder), "m.jpg", "8909735440"),
        folder=str(folder),
        images={"main": ["m.jpg"], "material": ["m-素材图.jpg"]},
        shop="A hat tshirt",
        template_spu="8909735440",
        main_image_index=1,
        title_cn="标题",
        title_en="Title",
        product_id=product_id,
    )


def step(name: str, status: Status, detail: str = "") -> StepResult:
    return StepResult(name=name, status=status, detail=detail)


# ------------------------------------------------------------------ 该写什么
def test_success_writes_product_id_and_status(tmp_path):
    update = updates_for(
        make_row(tmp_path),
        product_id="220114595101444742",
        failed=False,
        steps=[step("APPLY", Status.OK, "已提交")],
        statuses=STATUSES,
    )

    assert update.row_index == 2
    assert update.fields[COL_PRODUCT_ID] == "220114595101444742"
    assert update.fields[COL_STATUS] == "成功"
    assert update.fields[COL_PROBLEM] == "", "成功要把旧的问题项清掉"


def test_failure_without_product_id_leaves_it_alone(tmp_path):
    update = updates_for(
        make_row(tmp_path),
        product_id="",
        failed=True,
        steps=[step("PRECHECK", Status.FAIL, "没有「素材图」")],
        statuses=STATUSES,
    )

    assert COL_PRODUCT_ID not in update.fields, "没拿到ID 就不能乱写第 13 列"
    assert update.fields[COL_STATUS] == "失败"
    assert "没有「素材图」" in update.fields[COL_PROBLEM]


def test_failure_keeps_the_product_id_if_copy_succeeded(tmp_path):
    """复制成功了才失败的（比如换图写不进去）——ID 必须记下来，不然就是重复铺货。"""
    update = updates_for(
        make_row(tmp_path),
        product_id="220114595101444850",
        failed=True,
        steps=[step("APPLY", Status.FAIL, "图片没换上去")],
        statuses=STATUSES,
    )

    assert update.fields[COL_PRODUCT_ID] == "220114595101444850"
    assert update.fields[COL_STATUS] == "失败"


def test_problem_is_taken_from_the_first_failed_step(tmp_path):
    steps = [
        step("PRECHECK", Status.OK, "没问题"),
        step("COPY", Status.FAIL, "复制被拒"),
        step("APPLY", Status.FAIL, "后面还有别的错"),
    ]
    assert first_problem(True, steps).startswith("[COPY] 复制被拒")
    assert first_problem(False, steps) == ""


def test_long_problem_is_clipped():
    detail = "很长的原因" * 200
    got = first_problem(True, [step("COPY", Status.FAIL, detail)])
    assert len(got) <= PROBLEM_MAX
    assert got.endswith("…")


def test_statuses_fall_back_when_config_lacks_them():
    from temu_bot.sheet.writeback import statuses_of

    class Bare:
        statuses: dict[str, str] = {}

    assert statuses_of(None, Bare())["success"] == "成功"  # type: ignore[arg-type]


# ------------------------------------------------------------------ 真写
def test_write_back_updates_the_table(sandbox):
    cfg = sandbox["cfg"]
    table = sandbox["table"]
    row = make_row(sandbox["tmp"])

    report = write_back(
        cfg,
        updates_for(
            row, product_id="220114595101444742", failed=False, steps=[], statuses=STATUSES
        ),
    )

    assert report.ok, report.message
    assert report.backup and report.backup.exists(), "写前必须先备份"
    ws = openpyxl.load_workbook(table).active
    assert ws.cell(row=2, column=COL_PRODUCT_ID).value == "220114595101444742"
    assert ws.cell(row=2, column=COL_STATUS).value == "成功"


def test_write_back_reports_busy_instead_of_raising(sandbox, monkeypatch):
    """表被占用 → 返回 ok=False，**不抛异常**（上架结果要保住）。"""
    import temu_bot.sheet.writeback as wb
    from temu_bot.sheet.guard import TableBusyError

    def boom(*_a, **_k):
        raise TableBusyError("表正被 Excel/WPS 打开")

    monkeypatch.setattr(wb, "write_updates", boom)
    report = write_back(cfg=sandbox["cfg"], update=updates_for(
        make_row(sandbox["tmp"]), product_id="1", failed=False, steps=[], statuses=STATUSES
    ))

    assert report.ok is False
    assert "Excel" in report.message


# ------------------------------------------------------------------ 流程里的这一步
@pytest.fixture()
def runnable(sandbox, tmp_path):
    """一张"真能跑"的沙箱表：第 2 行的文件夹与图片都真的存在。

    上面那些表行指向 `C:\\tmp\\shop_a`（不存在），PRECHECK 就会挂，
    测不到 WRITEBACK 本身。所以这里另造一张。
    """
    from PIL import Image

    from tests.conftest import make_table
    from temu_bot.config import ShopCfg

    cfg = sandbox["cfg"]
    cfg.shops = [
        ShopCfg(name="A hat tshirt", login="hat-tshirt", expected_dxm_shops=["A hat tshirt"])
    ]

    folder = tmp_path / "imgs"
    folder.mkdir()
    Image.new("RGB", (1340, 1785), (10, 20, 30)).save(folder / "m.jpg", quality=80)
    Image.new("RGB", (800, 800), (30, 10, 20)).save(folder / "m-素材图.jpg", quality=80)

    table = make_table(
        tmp_path / "可跑表.xlsx",
        [
            {
                1: str(folder),
                2: "m.jpg",
                8: "m-素材图.jpg",
                9: "m",
                10: "A hat tshirt",
                11: "8909735440",
                12: 1,
                17: "新标题",
                18: "New Title",
            }
        ],
    )
    # 回写按 `cfg.project.table` 落盘，所以配置也得指到这张新表上
    cfg.project.table = table
    row = next(iter(SheetTable(table, "Sheet1", cfg.sheet).iter_rows()))
    return {"cfg": cfg, "table": table, "row": row}


def _pipeline(cfg, row, client, **kwargs):
    return ListingPipeline(cfg, row, login="hat-tshirt", client_factory=lambda: client, **kwargs)


def test_dry_run_never_writes_the_table(runnable):
    """干跑：WRITEBACK 必须是 🔸，表格一个格子都不许动。"""
    from tests.test_listing import FakeClient, product_with_images

    before = runnable["table"].read_bytes()
    client = FakeClient(
        template_row={"id": 111, "productName": "模板商品", "dxmState": "online"},
        product=product_with_images(),
    )
    outcome = _pipeline(runnable["cfg"], runnable["row"], client, live=False).run()

    step = next(s for s in outcome.steps if s.name == "WRITEBACK")
    assert step.status is Status.DRY
    assert outcome.writeback_ok is None
    assert runnable["table"].read_bytes() == before


def test_cancelled_row_does_not_touch_the_table(runnable):
    """还没动写操作就被取消 → 什么也没发生，不该往表里写状态。"""
    from tests.test_listing import FakeClient

    before = runnable["table"].read_bytes()
    outcome = _pipeline(
        runnable["cfg"], runnable["row"], FakeClient(), live=True, should_cancel=lambda: True
    ).run()

    assert outcome.writeback_ok is None
    step = next(s for s in outcome.steps if s.name == "WRITEBACK")
    assert step.status is Status.SKIP
    assert runnable["table"].read_bytes() == before


def test_successful_row_writes_id_and_status(runnable):
    """跑成功后，商品ID 与状态必须落到表里（这就是防重复铺货的那把锁）。"""
    from tests.test_listing import FakeClient, product_with_images

    client = FakeClient(
        template_row={"id": 111, "productName": "模板商品", "dxmState": "online"},
        product=product_with_images(),
        offline_rows=[{"id": 900, "productName": "别的商品"}],
        copy_adds={"id": 220114595101444742, "productName": "模板商品"},
    )
    outcome = _pipeline(runnable["cfg"], runnable["row"], client, live=True, op=1).run()

    assert not outcome.failed, [s.line() for s in outcome.steps]
    assert outcome.writeback_ok is True
    ws = openpyxl.load_workbook(runnable["table"]).active
    assert ws.cell(row=2, column=COL_PRODUCT_ID).value == "220114595101444742"
    assert ws.cell(row=2, column=COL_STATUS).value == "成功"


def test_failed_row_still_gets_its_status_written(runnable):
    """失败的行也要写状态——不然表里一直空着，下次还当"新建"跑。"""
    from tests.test_listing import FakeClient

    outcome = _pipeline(
        runnable["cfg"], runnable["row"], FakeClient(template_row=None), live=True
    ).run()

    assert outcome.failed
    assert outcome.writeback_ok is True
    ws = openpyxl.load_workbook(runnable["table"]).active
    assert ws.cell(row=2, column=COL_STATUS).value == "失败"
    assert "模板" in str(ws.cell(row=2, column=COL_PROBLEM).value)
    assert ws.cell(row=2, column=COL_PRODUCT_ID).value in (None, ""), "没拿到ID 就别写"


def test_writeback_failure_does_not_mark_the_listing_failed(runnable, monkeypatch):
    """回写失败**不算上架失败**——商品已经建好了，说"失败"会诱使人重跑（→重复铺货）。"""
    import temu_bot.sheet.writeback as wb
    from temu_bot.sheet.guard import TableBusyError
    from tests.test_listing import FakeClient, product_with_images

    monkeypatch.setattr(
        wb, "write_updates", lambda *a, **k: (_ for _ in ()).throw(TableBusyError("表开着"))
    )
    client = FakeClient(
        template_row={"id": 111, "productName": "模板商品", "dxmState": "online"},
        product=product_with_images(),
        offline_rows=[{"id": 900, "productName": "别的商品"}],
        copy_adds={"id": 220114595101444742, "productName": "模板商品"},
    )
    outcome = _pipeline(runnable["cfg"], runnable["row"], client, live=True, op=1).run()

    assert outcome.writeback_ok is False
    assert outcome.writeback_failed
    assert not outcome.failed, "上架是成功的，不能说成失败"
    step = next(s for s in outcome.steps if s.name == "WRITEBACK")
    assert step.status is Status.FAIL
    assert "table writeback" in step.detail, "要给出补写命令"


def test_custom_writeback_callback_is_used(runnable):
    """批量跑时上层会注入自己的回写（一整批只备份一次）。"""
    from tests.test_listing import FakeClient

    seen: list[int] = []

    class Report:
        ok = True
        message = "假装写了"

    def on_writeback(_pipeline, update):
        seen.append(update.row_index)
        return Report()

    outcome = _pipeline(
        runnable["cfg"],
        runnable["row"],
        FakeClient(template_row=None),
        live=True,
        on_writeback=on_writeback,
    ).run()

    assert seen == [runnable["row"].row_index]
    step = next(s for s in outcome.steps if s.name == "WRITEBACK")
    assert step.status is Status.OK
    assert "假装写了" in step.detail


# ------------------------------------------------------------------ 记账
def test_run_log_is_json_safe(runnable):
    """结果里出现非 JSON 类型也不能让日志写崩（接口/日志都靠它）。"""
    from tests.test_listing import FakeClient

    outcome = _pipeline(
        runnable["cfg"], runnable["row"], FakeClient(template_row=None), live=True
    ).run()
    payload = json.dumps(outcome.as_dict(), ensure_ascii=False)

    assert "writeback_ok" in payload
    assert json.loads(payload)["failed"] is True
