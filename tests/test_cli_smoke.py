"""CLI 冒烟测试。

教训：`cli.py` 里的 import 是**写在函数体里**的（为了启动速度），
所以 `--help` 不会触发它——命令真跑起来才发现 `ImportError`。
这里用 CliRunner 把每个命令至少执行一次，把这类错挡在提交前。
"""

from __future__ import annotations

import pytest
from typer.testing import CliRunner

from temu_bot.cli import app

runner = CliRunner()


@pytest.fixture()
def isolated(tmp_path, monkeypatch):
    """把配置**和项目根**都指到沙箱，别碰用户的真实商品表与 data/。

    只设 `TEMU_BOT_CONFIG` 不够：那只会换配置文件，`data/`（备份、日志、state）
    还算在真实项目根下，命令一写文件就落到真目录里了。
    """
    from tests.conftest import make_table

    table = make_table(tmp_path / "表.xlsx", [])
    config = tmp_path / "config.yaml"
    config.write_text(
        "\n".join(
            [
                "project:",
                f"  table: {table}",
                "  sheet_name: Sheet1",
                "toolbox:",
                f"  app_dir: {tmp_path / 'tb'}",
                f"  ledger: {tmp_path / 'tb' / 'used.txt'}",
            ]
        ),
        encoding="utf-8",
    )
    monkeypatch.setenv("TEMU_BOT_CONFIG", str(config))
    monkeypatch.setenv("TEMU_BOT_ROOT", str(tmp_path))
    return tmp_path


def test_help_lists_key_commands():
    result = runner.invoke(app, ["--help"])
    assert result.exit_code == 0
    for name in ("run", "preflight", "images", "doctor", "dxm", "table", "serve", "toolbox"):
        assert name in result.stdout


def test_serve_help_is_available():
    """serve 的导入写在函数体里，--help 不触发；但至少确认命令注册上了。"""
    result = runner.invoke(app, ["serve", "--help"])
    assert result.exit_code == 0
    assert "--port" in result.stdout
    assert "--no-open" in result.stdout


def test_serve_refuses_non_local_host_immediately():
    """只允许本机绑定——这条要在起服务之前就拦住。"""
    from temu_bot.web.app import serve as serve_app

    with pytest.raises(ValueError):
        serve_app(host="0.0.0.0", port=8899, open_browser=False)


def test_project_root_follows_the_exe_when_frozen(tmp_path, monkeypatch):
    """打包成 exe 后，配置/数据目录必须落在 exe 旁边，而不是被封进包里。"""
    import sys

    from temu_bot.config import project_root

    fake_exe = tmp_path / "temu-bot.exe"
    fake_exe.write_bytes(b"x")
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "executable", str(fake_exe))

    assert project_root() == tmp_path


def test_project_root_follows_the_source_tree_when_not_frozen():
    from temu_bot.config import project_root

    assert (project_root() / "src" / "temu_bot" / "config.py").exists()


def test_images_subcommands_run(isolated):
    """images 的每个子命令真的能被调用（光 --help 抓不到函数体内的 ImportError）。"""
    assert runner.invoke(app, ["images", "specs"]).exit_code == 0

    check = runner.invoke(app, ["images", "check", "--row-index", "2"])
    assert "ImportError" not in str(check.exception or ""), check.exception

    prep = runner.invoke(app, ["images", "prep", "--row-index", "2"])
    assert "ImportError" not in str(prep.exception or ""), prep.exception


def test_toolbox_subcommands_run(isolated):
    """toolbox 的每个子命令真的能跑（函数体内的 import 光靠 --help 抓不到）。"""
    from PIL import Image

    folder = isolated / "素材"
    folder.mkdir()
    Image.new("RGB", (20, 20), (1, 2, 3)).save(folder / "裙子.jpg")
    Image.new("RGB", (20, 20), (3, 2, 1)).save(folder / "帽子.png")

    ids_cmd = runner.invoke(app, ["toolbox", "ids"])
    assert ids_cmd.exit_code == 0, ids_cmd.exception
    assert "台账" in ids_cmd.stdout

    for argv in (
        ["toolbox", "rename", "--dir", str(folder), "--mode", "id-prefix"],
        ["toolbox", "export", "--dir", str(folder)],
        ["toolbox", "convert", "--dir", str(folder), "--to", "webp"],
        ["toolbox", "merge", "--dir", str(folder)],
    ):
        result = runner.invoke(app, argv)
        assert result.exit_code == 0, (argv, result.exception)
        assert "ImportError" not in str(result.exception or ""), result.exception
        assert "干跑" in result.stdout, argv

    # 干跑一个文件都不该动
    assert sorted(p.name for p in folder.iterdir()) == ["帽子.png", "裙子.jpg"]


def test_preflight_runs_on_empty_table(isolated):
    result = runner.invoke(app, ["preflight"])
    # 空表 → 校验不通过，但必须是"报出来"而不是崩掉
    assert result.exit_code == 1
    assert "ImportError" not in str(result.exception or ""), result.exception
    assert "上架前输入校验" in result.stdout


def _writeback_setup(tmp_path, monkeypatch, run_lines: list[str]):
    """造一张表 + 一个运行日志，把配置指过去。"""
    from tests.conftest import make_table

    table = make_table(
        tmp_path / "表.xlsx",
        [
            # 这一行是"新建"（没有商品ID）
            {1: str(tmp_path), 2: "a.jpg", 8: "a-素材图.jpg", 10: "A hat tshirt",
             11: "9082676134", 17: "标题", 18: "Title"},
            # 这一行表里**已经有**商品ID，且与日志里的不一样
            {1: str(tmp_path), 2: "b.jpg", 8: "b-素材图.jpg", 10: "A hat tshirt",
             11: "9082676134", 13: "999999999", 17: "标题B", 18: "TitleB"},
        ],
    )
    config = tmp_path / "config.yaml"
    config.write_text(
        "\n".join(
            [
                "project:",
                f"  table: {table}",
                "  sheet_name: Sheet1",
                "toolbox:",
                f"  app_dir: {tmp_path / 'tb'}",
                f"  ledger: {tmp_path / 'tb' / 'used.txt'}",
            ]
        ),
        encoding="utf-8",
    )
    monkeypatch.setenv("TEMU_BOT_CONFIG", str(config))
    # 备份写到 data/ 下，所以项目根也得挪进沙箱——否则会往真实 data/backups 里丢文件
    monkeypatch.setenv("TEMU_BOT_ROOT", str(tmp_path))

    log = tmp_path / "live_runs.jsonl"
    log.write_text("\n".join(run_lines) + "\n", encoding="utf-8")
    return table, log


def test_writeback_dry_run_does_not_touch_the_table(tmp_path, monkeypatch):
    """`table writeback` 默认只报告；不加 --live 一个格子都不许动。"""
    import openpyxl

    table, log = _writeback_setup(
        tmp_path,
        monkeypatch,
        ['{"row_index": 2, "ok": true, "product_id": "220114595101444742"}'],
    )
    before = table.read_bytes()

    result = runner.invoke(app, ["table", "writeback", "--from", str(log)])
    assert result.exit_code == 0, result.exception
    assert "220114595101444742" in result.stdout
    assert "干跑" in result.stdout
    assert table.read_bytes() == before

    wb = openpyxl.load_workbook(table)
    assert wb.active.cell(row=2, column=13).value in (None, "")


def test_writeback_fills_product_id_and_status(tmp_path, monkeypatch):
    import openpyxl

    table, log = _writeback_setup(
        tmp_path,
        monkeypatch,
        [
            '{"row_index": 2, "ok": true, "product_id": "220114595101444742"}',
            # 失败的行不该写
            '{"row_index": 3, "ok": false, "product_id": ""}',
        ],
    )

    result = runner.invoke(app, ["table", "writeback", "--from", str(log), "--live"])
    assert result.exit_code == 0, result.exception

    wb = openpyxl.load_workbook(table)
    ws = wb.active
    assert ws.cell(row=2, column=13).value == "220114595101444742"
    assert ws.cell(row=2, column=14).value == "成功"
    # 失败的那行、以及表里已有别的商品ID 的那行，都不能被动
    assert ws.cell(row=3, column=13).value == "999999999"


def test_writeback_refuses_to_overwrite_a_different_product_id(tmp_path, monkeypatch):
    """表里的商品ID 和日志里的不一样 → 绝不覆盖（那是两码事，得人工看）。"""
    import openpyxl

    table, log = _writeback_setup(
        tmp_path,
        monkeypatch,
        [
            '{"row_index": 3, "ok": true, "product_id": "220114595101444742"}',
        ],
    )

    result = runner.invoke(app, ["table", "writeback", "--from", str(log), "--live"])
    assert result.exit_code == 0, result.exception
    assert "表里已有商品ID" in result.stdout

    wb = openpyxl.load_workbook(table)
    assert wb.active.cell(row=3, column=13).value == "999999999", "不能被日志里的值盖掉"


def test_writeback_backs_up_before_writing(tmp_path, monkeypatch):
    table, log = _writeback_setup(
        tmp_path,
        monkeypatch,
        ['{"row_index": 2, "ok": true, "product_id": "220114595101444742"}'],
    )

    result = runner.invoke(app, ["table", "writeback", "--from", str(log), "--live"])
    assert result.exit_code == 0, result.exception

    backups = list((tmp_path / "data" / "backups").glob("*.bak.xlsx"))
    assert backups, "写前必须先备份"
    assert "备份" in result.stdout


def test_swap_images_flag_is_tri_state():
    """`--swap-images` 必须是三态：不给 = None（按行自动判断），给了才是 True。

    踩过的坑：一开始写成 `bool = Option(False, ...)`，于是"没给"和"明确不要"
    分不开，**新行也被当成已有商品**、不传图了。这条把那个回归钉住。
    """
    import inspect

    from temu_bot import cli

    sig = inspect.signature(cli.run)
    param = sig.parameters["swap_images"]
    assert param.default.default is None, "默认必须是 None（=自动），不能是 False"

