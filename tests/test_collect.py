"""收集图片测试：配对规则、缺素材图跳过、绝不覆盖已有表。"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from PIL import Image

from temu_bot.collect import (
    CollectOptions,
    default_output,
    safe_output_path,
    scan_folder,
    write_table,
)
from temu_bot.sheet import schema

MAIN_SIZE = (1340, 1785)
MAT_SIZE = (800, 800)


def img(path: Path, size=MAIN_SIZE) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, (30, 90, 160)).save(path, quality=85)
    return path


@pytest.fixture()
def src(tmp_path: Path) -> Path:
    """一个典型素材目录：3 对完整的 + 1 张缺素材图 + 1 张多余的素材图。"""
    root = tmp_path / "素材"
    root.mkdir()
    img(root / "00742093597.jpg")
    img(root / "00742093597-素材图.jpg", MAT_SIZE)
    img(root / "01409714133.jpg")
    img(root / "01409714133-素材图.jpg", MAT_SIZE)
    img(root / "01564109050.jpg")
    img(root / "01564109050-素材图.jpg", MAT_SIZE)
    img(root / "09999999999.jpg")            # ← 没有素材图，应被跳过
    img(root / "05000000000-素材图.jpg", MAT_SIZE)   # ← 没有主图，不该被当成主图
    (root / "说明.txt").write_text("不是图片", encoding="utf-8")
    return root


def opts(folder: Path, **kw) -> CollectOptions:
    base = {"folder": folder, "shop": "A lovely costume", "template_spu": "8909735440"}
    base.update(kw)
    return CollectOptions(**base)


# ------------------------------------------------------------------ 配对
def test_pairs_main_with_material(src: Path):
    result = scan_folder(opts(src))

    assert result.image_count == 8          # 图片数（不含 txt）
    assert [r.main_name for r in result.rows] == [
        "00742093597.jpg",
        "01409714133.jpg",
        "01564109050.jpg",
    ]
    assert result.rows[0].material_name == "00742093597-素材图.jpg"
    assert result.rows[0].stem == "00742093597"


def test_row_without_material_is_skipped(src: Path):
    """用户明确要求：缺素材图就跳过该行。"""
    result = scan_folder(opts(src))

    assert "09999999999.jpg" not in [r.main_name for r in result.rows]
    skips = {s.name: s.reason for s in result.skipped}
    assert "09999999999.jpg" in skips
    assert "素材图" in skips["09999999999.jpg"]


def test_material_only_file_is_not_treated_as_main(src: Path):
    """只有素材图、没有主图的，不能当成一行商品。"""
    result = scan_folder(opts(src))
    assert "05000000000-素材图.jpg" not in [r.main_name for r in result.rows]


def test_rows_are_sorted_by_name(src: Path):
    names = [r.main_name for r in scan_folder(opts(src)).rows]
    assert names == sorted(names, key=str.lower)


def test_missing_folder_is_reported(tmp_path: Path):
    result = scan_folder(opts(tmp_path / "没有这个目录"))
    assert not result.ok
    assert "不存在" in result.skipped[0].reason


def test_empty_folder_is_reported(tmp_path: Path):
    empty = tmp_path / "空的"
    empty.mkdir()
    result = scan_folder(opts(empty))
    assert not result.ok
    assert "一张图都没有" in result.skipped[0].reason


def test_extension_mismatch_still_pairs(tmp_path: Path):
    """主图 .jpg、素材图 .png —— 名字对得上就算一对。"""
    root = tmp_path / "素材2"
    root.mkdir()
    img(root / "123.jpg")
    img(root / "123-素材图.png", MAT_SIZE)

    result = scan_folder(opts(root))
    assert len(result.rows) == 1
    assert result.rows[0].material_name == "123-素材图.png"


def test_recursive_collect_can_be_turned_off(tmp_path: Path):
    """默认只看所选文件夹这一层——父目录里套着大子树时，递归会慢很多。"""
    root = tmp_path / "父目录"
    (root / "子目录").mkdir(parents=True)
    img(root / "top.jpg")
    img(root / "top-素材图.jpg", MAT_SIZE)
    img(root / "子目录" / "inner.jpg")
    img(root / "子目录" / "inner-素材图.jpg", MAT_SIZE)

    shallow = scan_folder(opts(root))
    assert [r.main_name for r in shallow.rows] == ["top.jpg"]

    deep = scan_folder(opts(root, recursive=True))
    assert sorted(r.main_name for r in deep.rows) == ["inner.jpg", "top.jpg"]


def test_summary_says_which_collection_type_was_used(tmp_path: Path):
    root = tmp_path / "父目录"
    root.mkdir()
    img(root / "a.jpg")

    assert "含子文件夹" not in scan_folder(opts(root)).summary()
    assert "含子文件夹" in scan_folder(opts(root, recursive=True)).summary()


def test_scan_keeps_the_index_for_reuse(src: Path):
    """索引要带回去给调用方复用，否则查主图规格会把整个目录再扫一遍。"""
    result = scan_folder(opts(src))
    assert result.index is not None
    assert len(result.index.files) == result.image_count


def test_force_generate_marks_rows_instead_of_skipping(src: Path):
    result = scan_folder(opts(src, force_gen_material=True))
    names = [r.main_name for r in result.rows]
    assert "09999999999.jpg" in names
    row = next(r for r in result.rows if r.main_name == "09999999999.jpg")
    assert row.generated_material is True
    assert row.material_name == "09999999999-素材图.jpg"


def test_summary_is_readable(src: Path):
    s = scan_folder(opts(src)).summary()
    assert "张图" in s and "行" in s and "跳过" in s


# ------------------------------------------------------------------ 写表
def test_write_produces_the_expected_columns(src: Path):
    result = scan_folder(opts(src))
    out = write_table(result, src / "新表.xlsx")

    wb = openpyxl.load_workbook(out.path, read_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = list(rows[0])
    expected = [schema.EXPECTED_HEADER[i] for i in sorted(schema.EXPECTED_HEADER)]
    assert header == expected
    assert len(header) == 20

    first = rows[1]
    def cell(col):
        return first[col - 1]

    assert cell(schema.COL_FOLDER) == str(src)
    assert cell(schema.COL_MAIN_IMAGE) == "00742093597.jpg"
    assert cell(schema.COL_MATERIAL_IMAGE) == "00742093597-素材图.jpg"
    assert cell(schema.COL_SKU_PREFIX_CODE) == "00742093597"
    assert cell(schema.COL_SHOP) == "A lovely costume"
    assert cell(schema.COL_TEMPLATE_SPU) == "8909735440"
    assert cell(schema.COL_MAIN_IMAGE_INDEX) == 1
    # 这几列留给上架阶段回填，收集时必须是空的
    for col in (schema.COL_PRODUCT_ID, schema.COL_STATUS, schema.COL_TITLE_CN,
                schema.COL_TITLE_EN, schema.COL_SKC_CODE):
        assert cell(col) in ("", None), f"第 {col} 列不该有值"


def test_use_img_name_as_skc_fills_column_19(src: Path):
    result = scan_folder(opts(src, use_img_name_as_skc=True))
    out = write_table(result, src / "带SKC.xlsx")

    wb = openpyxl.load_workbook(out.path, read_only=True)
    rows = list(wb["Sheet1"].iter_rows(values_only=True))
    wb.close()
    assert rows[1][schema.COL_SKC_CODE - 1] == "00742093597"


def test_never_overwrites_an_existing_table(src: Path):
    """素材目录里本来就躺着一张表——绝不能覆盖。"""
    existing = src / "商品信息表.xlsx"
    existing.write_bytes("这是用户原有的表，不许动".encode("utf-8"))
    before = existing.read_bytes()

    result = scan_folder(opts(src))
    out = write_table(result)          # 不指定路径 → 默认就是上面那个名字

    assert out.renamed is True
    assert out.path != existing
    assert existing.read_bytes() == before, "原表被改了！"
    assert out.path.exists()


def test_safe_output_path_keeps_original_when_free(tmp_path: Path):
    target = tmp_path / "新表.xlsx"
    final, renamed = safe_output_path(target)
    assert final == target
    assert renamed is False


def test_generated_material_is_actually_written(src: Path):
    result = scan_folder(opts(src, force_gen_material=True))
    assert not (src / "09999999999-素材图.jpg").exists()

    out = write_table(result, src / "强生成.xlsx", generate_missing_material=True)

    made = src / "09999999999-素材图.jpg"
    assert made.exists(), "勾了强制生成就该真的生成出来"
    assert made.name in out.generated_materials
    with Image.open(made) as im:
        assert im.size == MAT_SIZE


def test_write_without_rows_is_rejected(tmp_path: Path):
    empty = tmp_path / "空目录"
    empty.mkdir()
    with pytest.raises(ValueError):
        write_table(scan_folder(opts(empty)), empty / "x.xlsx")


def test_default_output_lives_in_the_folder(src: Path):
    assert default_output(src) == src / "商品信息表.xlsx"


# ------------------------------------------------------------------ 主图规格预检
def test_check_main_images_flags_bad_sizes(tmp_path: Path):
    """收集时提前发现主图尺寸不合规，别等上架被平台打回。"""
    from temu_bot.collect import check_main_images

    root = tmp_path / "素材"
    root.mkdir()
    img(root / "ok.jpg", (1340, 1785))
    img(root / "ok-素材图.jpg", MAT_SIZE)
    img(root / "short.jpg", (1340, 1784))          # 矮 1 像素，就是真实踩过的坑
    img(root / "short-素材图.jpg", MAT_SIZE)

    problems = check_main_images(scan_folder(opts(root)))

    assert "short.jpg" in problems
    assert any("高" in p for p in problems["short.jpg"])
    assert "ok.jpg" not in problems


def test_check_main_images_is_empty_when_all_good(tmp_path: Path):
    from temu_bot.collect import check_main_images

    root = tmp_path / "素材"
    root.mkdir()
    img(root / "a.jpg", (1340, 1785))
    img(root / "a-素材图.jpg", MAT_SIZE)

    assert check_main_images(scan_folder(opts(root))) == {}


# ------------------------------------------------------------------ 图片文件名清单 → 中/英文标题
def make_title_list(path: Path, rows: list[tuple[str, str, str]], *, english_col: int = 10) -> Path:
    """造一份和用户实际文件同构的清单：A=ID、B=中文、J=英文。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "图片文件名清单"
    ws.append(["数字前缀", "中文后缀 "])
    for key, cn, en in rows:
        # C..I 留空，英文放在 J 列（用户实际就是这样）
        ws.append([key, cn] + [""] * (english_col - 3) + [en])
    wb.save(path)
    return path


def test_titles_filled_by_11_digit_id(src: Path, tmp_path: Path):
    lst = make_title_list(
        tmp_path / "图片文件名清单.xlsx",
        [
            ("00742093597", "男士休闲图案印花短袖T恤夏季", "Men's T-Shirt Casual Graphic Print"),
            ("01409714133", "男士短袖圆领T恤万圣节南瓜", "Men's Halloween Pumpkin T-Shirt"),
        ],
    )
    result = scan_folder(opts(src, title_list=lst))
    by_stem = {r.sku_prefix: r for r in result.rows}

    assert by_stem["00742093597"].title_cn == "男士休闲图案印花短袖T恤夏季"
    assert by_stem["00742093597"].title_en == "Men's T-Shirt Casual Graphic Print"
    assert by_stem["01409714133"].title_en == "Men's Halloween Pumpkin T-Shirt"
    assert by_stem["01564109050"].title_cn == ""      # 清单里没有这个 ID
    assert result.titled_rows == 2


def test_titles_written_to_columns_17_18(src: Path, tmp_path: Path):
    lst = make_title_list(
        tmp_path / "清单.xlsx",
        [("00742093597", "中文标题", "Men's T-Shirt Casual Graphic Print")],
    )
    out = write_table(scan_folder(opts(src, title_list=lst)), tmp_path / "带标题.xlsx")

    ws = openpyxl.load_workbook(out.path).active
    row2 = [c.value for c in ws[2]]
    assert row2[schema.COL_TITLE_CN - 1] == "中文标题"
    assert row2[schema.COL_TITLE_EN - 1] == "Men's T-Shirt Casual Graphic Print"
    # 没匹配到的那几行留空
    row3 = [c.value for c in ws[3]]
    assert row3[schema.COL_TITLE_CN - 1] in (None, "")


def test_title_list_matches_leading_zeros(src: Path, tmp_path: Path):
    """ID 是文本 `00742093597`：绝不能转成 int（前导 0 会丢，就查不到了）。"""
    lst = make_title_list(tmp_path / "清单.xlsx", [("00742093597", "带前导零", "")])
    result = scan_folder(opts(src, title_list=lst))
    by_stem = {r.sku_prefix: r for r in result.rows}
    assert by_stem["00742093597"].title_cn == "带前导零"


def test_title_list_reads_gbk_csv(src: Path, tmp_path: Path):
    path = tmp_path / "图片文件名清单.csv"
    path.write_bytes(
        '"数字前缀","中文后缀 " \r\n"00742093597","中文标题" \r\n'.encode("gbk")
    )
    result = scan_folder(opts(src, title_list=path))
    by_stem = {r.sku_prefix: r for r in result.rows}
    assert by_stem["00742093597"].title_cn == "中文标题"
    assert by_stem["00742093597"].title_en == ""      # CSV 没有英文列


def test_english_column_shifted_up_one_row_is_realigned(src: Path, tmp_path: Path):
    """用户那份清单：A1 是表头，J1 却已经是一条英文标题 → 英文整列错一行。

    不处理的话，每个商品都会拿到**下一个商品**的英文。这里按顺序重新对齐。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["数字前缀", "中文后缀 "] + [""] * 7 + ["English for A"])
    ws.append(["00742093597", "中文一"] + [""] * 7 + ["English for B"])
    ws.append(["01409714133", "中文二"] + [""] * 7 + [""])
    lst = tmp_path / "错一行.xlsx"
    wb.save(lst)

    result = scan_folder(opts(src, title_list=lst))
    by_stem = {r.sku_prefix: r for r in result.rows}

    assert result.title_list is not None and result.title_list.shifted is True
    assert by_stem["00742093597"].title_en == "English for A"
    assert by_stem["01409714133"].title_en == "English for B"


def test_missing_title_list_is_reported_not_fatal(src: Path, tmp_path: Path):
    result = scan_folder(opts(src, title_list=tmp_path / "没有这个文件.xlsx"))
    assert result.title_list is not None
    assert result.title_list.error
    assert result.rows                              # 收集照常
    assert all(not r.title_cn for r in result.rows)


# ------------------------------------------------------------------ 店铺名核对
def test_find_unknown_shop_warns_on_a_typo():
    from temu_bot.collect import find_unknown_shop

    class S:
        name = "A lovely costume"

    class Cfg:
        shops = [S()]

    assert find_unknown_shop(Cfg(), "A lovely costume") == ""
    warn = find_unknown_shop(Cfg(), "A lovely costum")     # 少了一个 e
    assert "不在配置的店铺列表里" in warn
    assert "A lovely costume" in warn


def test_find_unknown_shop_warns_when_blank():
    from temu_bot.collect import find_unknown_shop

    class Cfg:
        shops = []

    assert "还没填店铺名" in find_unknown_shop(Cfg(), "  ")


def test_find_unknown_shop_is_silent_when_nothing_configured():
    """配置里一个店都没有时，不该乱报错。"""
    from temu_bot.collect import find_unknown_shop

    class Cfg:
        shops = []

    assert find_unknown_shop(Cfg(), "任意店名") == ""
