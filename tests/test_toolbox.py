"""工具箱四个工具的行为测试。

这些是**原生 Python 重写**，所以能像普通代码一样钉住行为——桌面那套只有
PowerShell 的 45 项自测，跑不进 pytest，改一行就得手工验。

重点钉的几条（都是桌面工具踩过坑之后才加上的）：
- 台账只增不删、发出的 ID 永不重发（还要合并桌面那份旧台账）
- 已经带 11 位前缀的文件不叠第二个 ID
- 导出 CSV 的字节格式（GBK / 无 BOM / 行尾一个空格）—— 下游流程认这个
- 干跑**一个字节都不写**
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from PIL import Image

from temu_bot.toolbox import convert, ids, listing, merge, rename
from temu_bot.toolbox.result import FAILED, OK, SKIPPED, WOULD


# ------------------------------------------------------------------ 小工具
def make_image(path: Path, size=(20, 20), mode="RGB", color=(10, 20, 30)) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new(mode, size, color).save(path)
    return path


def make_book(path: Path, sheets: dict[str, list[list[object]]]) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def names(folder: Path) -> list[str]:
    return sorted(p.name for p in folder.iterdir())


# ------------------------------------------------------------------ 台账
def test_new_id_is_eleven_digits_and_not_reused():
    taken: set[str] = set()
    got = {ids.new_unique_id(taken) for _ in range(50)}
    assert len(got) == 50, "50 次不该撞号"
    assert all(len(i) == 11 and i.isdigit() for i in got)
    assert taken == got, "发出去的号必须当场记进 taken"


def test_new_id_gives_up_instead_of_looping_forever(monkeypatch):
    """随机源坏了一直吐同一个号时，要**报错**而不是死循环。"""
    monkeypatch.setattr(ids.secrets, "token_bytes", lambda n: b"\x00" * n)
    taken = {"00000000000"}  # 这个号已经在台账里了
    with pytest.raises(ids.LedgerError):
        ids.new_unique_id(taken)


def test_ledger_merges_the_old_one_and_only_grows(sandbox):
    cfg = sandbox["cfg"]
    cfg.toolbox.ledger.parent.mkdir(parents=True, exist_ok=True)
    cfg.toolbox.ledger.write_text("11111111111\n22222222222\n", encoding="utf-8")

    assert ids.load_taken(cfg) == {"11111111111", "22222222222"}, "旧台账要能读进来"

    path = ids.save_ids(cfg, ["33333333333"])
    assert path == cfg.toolbox_ledger
    # 项目台账是新文件，只写新号；旧台账原封不动
    assert cfg.toolbox_ledger.read_text(encoding="utf-8") == "33333333333\n"
    assert ids.load_taken(cfg) == {"11111111111", "22222222222", "33333333333"}

    ids.save_ids(cfg, ["44444444444"])
    assert cfg.toolbox_ledger.read_text(encoding="utf-8") == "33333333333\n44444444444\n"


def test_ledger_file_has_no_bom(sandbox):
    cfg = sandbox["cfg"]
    ids.save_ids(cfg, ["12345678901"])
    assert cfg.toolbox_ledger.read_bytes()[:3] != b"\xef\xbb\xbf"


def test_save_ids_with_nothing_to_write_does_not_create_a_file(sandbox):
    cfg = sandbox["cfg"]
    assert ids.save_ids(cfg, []) is None
    assert not cfg.toolbox_ledger.exists()


# ------------------------------------------------------------------ SKU 重命名
@pytest.fixture()
def pics(tmp_path: Path) -> Path:
    folder = tmp_path / "素材"
    folder.mkdir()
    make_image(folder / "裙子.jpg")
    make_image(folder / "帽子.png")
    return folder


def test_id_only_replaces_the_whole_name(sandbox, pics):
    cfg = sandbox["cfg"]
    report = rename.run(cfg, pics, "id-only")

    assert report.ok
    after = names(pics)
    assert len(after) == 2
    assert all(re.fullmatch(r"\d{11}\.(jpg|png)", n) for n in after), after
    assert sorted(p.suffix for p in pics.iterdir()) == [".jpg", ".png"]


def test_id_prefix_keeps_the_old_name(sandbox, pics):
    cfg = sandbox["cfg"]
    rename.run(cfg, pics, "id-prefix")
    after = names(pics)
    assert "帽子.png" not in after
    assert any(re.fullmatch(r"\d{11}-帽子\.png", n) for n in after), after
    assert any(re.fullmatch(r"\d{11}-裙子\.jpg", n) for n in after), after


def test_already_prefixed_file_reuses_its_id_instead_of_stacking(sandbox, pics):
    """重复跑不能变成 `新ID-旧ID-原名` —— 那样导出的清单直接作废。"""
    cfg = sandbox["cfg"]
    make_image(pics / "12345678901-旧的.jpg")

    report = rename.run(cfg, pics, "id-prefix")
    after = names(pics)
    assert "12345678901-旧的.jpg" in after, after
    assert not any(n.count("-") > 1 for n in after), f"叠了第二个 ID：{after}"

    skip = [i for i in report.items if i.name == "12345678901-旧的.jpg"]
    assert skip and skip[0].status == SKIPPED and "沿用原 ID" in skip[0].detail
    # 沿用下来的 ID 也要进台账，免得以后发给别的图
    assert "12345678901" in ids.load_taken(cfg)


def test_strip_id_only_cuts_a_real_eleven_digit_prefix(sandbox):
    """前缀不是 11 位数字时**别乱裁** —— 原 .bat 是无条件按第一个 `-` 截断的。"""
    cfg = sandbox["cfg"]
    folder = sandbox["tmp"] / "还原"
    folder.mkdir()
    make_image(folder / "12345678901-裙子.jpg")
    make_image(folder / "abc-1.jpg")
    make_image(folder / "123-短前缀.jpg")
    make_image(folder / "没有横杠.jpg")

    report = rename.run(cfg, folder, "strip-id")
    after = names(folder)
    assert "12345678901.jpg" in after
    assert "abc-1.jpg" in after, "前缀不是 11 位数字，保持原样"
    assert "123-短前缀.jpg" in after
    assert "没有横杠.jpg" in after
    skipped = [i for i in report.items if i.status == SKIPPED]
    assert len(skipped) == 3


def test_roundtrip_ends_with_pure_ids_and_leaves_a_manifest(sandbox, pics):
    cfg = sandbox["cfg"]
    rename.run(cfg, pics, "roundtrip")

    after = names(pics)
    assert "图片文件名清单.csv" in after and "图片文件名清单.xlsx" in after
    image_names = [n for n in after if not n.startswith("图片文件名清单")]
    assert all(re.fullmatch(r"\d{11}\.(jpg|png)", n) for n in image_names), image_names


def test_roundtrip_reports_a_bad_export_format_instead_of_crashing(sandbox, pics):
    """两种清单格式都没勾 → 要报出来，而且**改名那部分照样跑完**。

    控制台把「SKU 重命名」和「导出图片清单」并成了一个动作（roundtrip），
    所以导出失败不能让整件事半途而废 —— 文件最后仍然要是纯 11 位 ID。
    """
    cfg = sandbox["cfg"]
    report = rename.run(cfg, pics, "roundtrip", export_format="")

    failed = [i for i in report.items if i.status == FAILED]
    assert failed and listing.NO_FORMAT in failed[0].detail
    image_names = [p.name for p in pics.iterdir()]
    assert all(re.fullmatch(r"\d{11}\.(jpg|png)", n) for n in image_names), image_names


def test_target_already_taken_is_skipped_not_overwritten(sandbox, pics):
    cfg = sandbox["cfg"]
    first = rename.run(cfg, pics, "id-only")
    before = names(pics)

    # 再来一次：新 ID 不会撞上，所以会正常改名 —— 这里换个构造方式：
    # 造一个"目标名已经被别的文件占着"的局面
    folder = sandbox["tmp"] / "撞名"
    folder.mkdir()
    make_image(folder / "a.jpg")
    taken_id = "77777777777"
    make_image(folder / f"{taken_id}.jpg")          # 这个名字已经被占
    monkeypatch_ids = iter([taken_id, "88888888888"])
    import temu_bot.toolbox.ids as ids_mod

    real = ids_mod.new_unique_id

    def fake_new(taken, length=11):
        candidate = next(monkeypatch_ids)
        taken.add(candidate)
        return candidate

    ids_mod.new_unique_id = fake_new
    try:
        report = rename.run(cfg, folder, "id-only")
    finally:
        ids_mod.new_unique_id = real

    skipped = [i for i in report.items if i.status == SKIPPED]
    assert skipped, report.summary()
    assert first.ok
    assert before  # 上面那批确实改了名


def test_plan_writes_nothing(sandbox, pics):
    cfg = sandbox["cfg"]
    before = names(pics)
    report = rename.plan(cfg, pics, "id-only")

    assert before == names(pics), "干跑改名了文件"
    assert report.counts[WOULD] == 2
    assert not cfg.toolbox_ledger.exists(), "干跑不该写台账"


def test_rename_on_a_folder_without_images_says_so(sandbox):
    cfg = sandbox["cfg"]
    empty = sandbox["tmp"] / "空"
    empty.mkdir()
    report = rename.plan(cfg, empty, "id-only")
    assert report.items == []
    assert "没有图片" in report.note


def test_rename_on_a_missing_folder_reports_instead_of_raising(sandbox):
    cfg = sandbox["cfg"]
    report = rename.plan(cfg, sandbox["tmp"] / "不存在", "id-only")
    assert not report.ok
    assert report.items[0].status == FAILED


# ------------------------------------------------------------------ 导出清单
def test_csv_is_byte_compatible_with_the_old_format(tmp_path: Path):
    """GBK、无 BOM、CRLF、**行尾一个空格** —— 下游流程认的就是这些字节。"""
    folder = tmp_path / "素材"
    folder.mkdir()
    for name in ("12345678901-裙子.jpg", "没有横杠.png", "带着-横杠-的名字.jpg"):
        make_image(folder / name)

    report = listing.run(folder, fmt="csv")
    assert report.ok, report.summary()

    raw = (folder / listing.CSV_NAME).read_bytes()
    assert raw[:3] != b"\xef\xbb\xbf", "不能有 BOM"
    text = raw.decode("gbk")
    lines = text.split("\r\n")
    # 按文件名排序：'1'(0x31) < '带' < '没'
    assert lines[0] == '"数字前缀","中文后缀" '
    assert lines[1] == '"12345678901","裙子" '
    assert lines[2] == '"带着","横杠-的名字" ', "只按第一个 - 拆，后面的横杠留在后缀里"
    assert lines[3] == '"没有横杠","" ', "没有 - 时整名当前缀、后缀空"
    assert lines[-1] == "", "应以 CRLF 收尾"
    assert text.endswith("\r\n")


def test_csv_quotes_inside_values_are_doubled(tmp_path: Path):
    """Windows 文件名里放不了 `"`，所以直接喂给写函数测。"""
    target = tmp_path / "清单.csv"
    listing.write_csv(target, [listing.Entry('说"引号"', '后缀"也是')])
    text = target.read_bytes().decode("gbk")
    assert text.split("\r\n")[1] == '"说""引号""","后缀""也是" '


def test_csv_writes_unencodable_characters_as_question_marks(tmp_path: Path):
    """文件名里有 GBK 表示不了的字符（emoji）→ 写成 ? 并继续，不中断。"""
    folder = tmp_path / "素材"
    folder.mkdir()
    make_image(folder / "AB-😀smile.jpg")
    report = listing.run(folder, fmt="csv")
    assert report.ok, report.summary()
    text = (folder / listing.CSV_NAME).read_bytes().decode("gbk", errors="replace")
    assert "?" in text


def test_xlsx_looks_like_the_old_one(tmp_path: Path):
    from openpyxl import load_workbook

    folder = tmp_path / "素材"
    folder.mkdir()
    make_image(folder / "12345678901-裙子.jpg")
    make_image(folder / "22222222222-帽子.jpg")
    listing.run(folder, fmt="xlsx")

    wb = load_workbook(folder / "图片文件名清单.xlsx")
    ws = wb[listing.XLSX_SHEET]
    assert ws["A1"].value == "数字前缀"
    assert ws["B1"].value == "中文后缀"
    assert ws["A1"].font.bold is True
    assert ws.freeze_panes == "A2", "冻结首行"
    assert ws.auto_filter.ref == "A1:B3", "首行要有筛选按钮"
    assert ws.column_dimensions["A"].width == 14
    assert ws.column_dimensions["B"].width == 80
    assert ws.max_row == 3
    wb.close()


def test_existing_list_is_backed_up_not_overwritten(tmp_path: Path):
    folder = tmp_path / "素材"
    folder.mkdir()
    make_image(folder / "12345678901-裙子.jpg")
    old = folder / listing.CSV_NAME
    old.write_bytes(b"old bytes")

    report = listing.run(folder, fmt="csv")
    assert report.ok, report.summary()
    backups = list(folder.glob(f"{listing.CSV_NAME}.bak-*"))
    assert len(backups) == 1, "旧清单必须先备份"
    assert backups[0].read_bytes() == b"old bytes"
    assert old.read_bytes() != b"old bytes"
    assert any(i.action == "备份旧清单" for i in report.items)


def test_export_plan_writes_nothing(tmp_path: Path):
    folder = tmp_path / "素材"
    folder.mkdir()
    make_image(folder / "12345678901-裙子.jpg")

    report = listing.plan(folder, fmt="both")
    assert report.counts[WOULD] == 2
    assert names(folder) == ["12345678901-裙子.jpg"]
    assert "目标已存在" not in report.items[0].detail


def test_export_refuses_when_no_format_is_chosen(tmp_path: Path):
    folder = tmp_path / "素材"
    folder.mkdir()
    report = listing.run(folder, fmt="")
    assert not report.ok
    assert listing.NO_FORMAT in report.items[0].detail


# ------------------------------------------------------------------ 格式转换
def test_convert_png_to_webp_and_avif(tmp_path: Path):
    folder = tmp_path / "图"
    folder.mkdir()
    make_image(folder / "a.png", size=(40, 30))

    report = convert.run(folder, target_format="webp")
    assert report.ok, report.summary()
    with Image.open(folder / "a.webp") as im:
        assert im.format == "WEBP"
        assert im.size == (40, 30)

    report = convert.run(folder, target_format="avif")
    assert report.ok, report.summary()
    with Image.open(folder / "a.avif") as im:
        assert im.format == "AVIF"
        assert im.size == (40, 30)


def test_convert_to_jpg_flattens_transparency(tmp_path: Path):
    folder = tmp_path / "图"
    folder.mkdir()
    make_image(folder / "透明.png", size=(10, 10), mode="RGBA", color=(255, 0, 0, 0))

    report = convert.run(folder, target_format="jpg")
    assert report.ok, report.summary()
    with Image.open(folder / "透明.jpg") as im:
        assert im.format == "JPEG"
        assert im.mode == "RGB"
        # 全透明 → 铺成白底，而不是黑块
        assert im.getpixel((5, 5)) == (255, 255, 255)


def test_convert_produces_the_expected_extension_names(tmp_path: Path):
    assert convert.EXT_BY_FORMAT["jpeg"] == "jpg"
    folder = tmp_path / "图"
    folder.mkdir()
    make_image(folder / "a.png")
    convert.run(folder, target_format="jpeg")
    assert (folder / "a.jpg").is_file()


def test_same_name_collisions_get_an_underscore_suffix(tmp_path: Path):
    """两个子文件夹里有同名图，一起转到一个输出目录时不能互相盖掉。"""
    src = tmp_path / "src"
    make_image(src / "one" / "同名.png")
    make_image(src / "two" / "同名.png")
    out = tmp_path / "out"

    report = convert.run(
        src, target_format="webp", output_dir=out, recursive=True, overwrite=True
    )
    assert report.ok, report.summary()
    assert sorted(p.name for p in out.iterdir()) == ["同名.webp", "同名_1.webp"]


def test_convert_skips_when_target_is_the_source_itself(tmp_path: Path):
    """同一个目录里既转又扫时，不能把自己刚转出来的 webp 再转一遍。"""
    folder = tmp_path / "图"
    folder.mkdir()
    make_image(folder / "a.png")

    assert convert.run(folder, target_format="webp").ok
    second = convert.run(folder, target_format="webp")

    assert sorted(p.name for p in folder.iterdir()) == ["a.png", "a.webp"], "多转了"
    assert not second.items or all(i.status != OK for i in second.items)


def test_convert_dry_run_writes_nothing_but_still_probes(tmp_path: Path):
    folder = tmp_path / "图"
    folder.mkdir()
    make_image(folder / "a.png", size=(40, 30))

    report = convert.plan(folder, target_format="webp")
    assert names(folder) == ["a.png"]
    assert report.items[0].status == WOULD
    assert "40x30" in report.items[0].detail


def test_convert_reports_a_broken_file_instead_of_crashing(tmp_path: Path):
    folder = tmp_path / "图"
    folder.mkdir()
    make_image(folder / "好的.png")
    (folder / "坏的.png").write_bytes(b"not an image at all")

    report = convert.run(folder, target_format="webp")
    bad = [i for i in report.items if i.status == FAILED]
    assert len(bad) == 1 and "打不开" in bad[0].detail
    assert (folder / "好的.webp").is_file(), "一个坏文件不该影响别的"


def test_convert_rejects_an_unknown_target_format(tmp_path: Path):
    folder = tmp_path / "图"
    folder.mkdir()
    report = convert.plan(folder, target_format="gif")
    assert not report.ok
    assert "不支持" in report.items[0].detail


def test_heic_is_flagged_with_a_clear_reason(tmp_path: Path):
    """Pillow 没有 HEIC 解码器 —— 报错要写明白，别让人以为是文件坏了。"""
    folder = tmp_path / "图"
    folder.mkdir()
    (folder / "照片.heic").write_bytes(b"\x00\x00\x00\x18ftypheic")

    report = convert.plan(folder, target_format="png")
    assert not report.ok
    assert "HEIC" in report.items[0].detail


# ------------------------------------------------------------------ Excel 合并
def test_merge_concatenates_every_sheet_in_order(tmp_path: Path):
    folder = tmp_path / "表"
    make_book(folder / "one.xlsx", {"S1": [["甲", "乙"], [1, 2]], "S2": [["丙"]]})
    make_book(folder / "two.xlsx", {"S1": [["丁", "戊", "己"]]})

    report = merge.run(folder)
    assert report.ok, report.summary()

    from openpyxl import load_workbook

    wb = load_workbook(folder / merge.DEFAULT_NAME)
    rows = [list(r) for r in wb[merge.SHEET_NAME].iter_rows(values_only=True)]
    wb.close()
    assert rows == [["甲", "乙", None], [1, 2, None], ["丙", None, None], ["丁", "戊", "己"]]


def test_merge_can_insert_a_source_row_per_workbook(tmp_path: Path):
    folder = tmp_path / "表"
    make_book(folder / "one.xlsx", {"S1": [["甲"], ["乙"]]})
    make_book(folder / "two.xlsx", {"S1": [["丙"]]})

    merge.run(folder, add_source_row=True)
    from openpyxl import load_workbook

    wb = load_workbook(folder / merge.DEFAULT_NAME)
    rows = [list(r) for r in wb[merge.SHEET_NAME].iter_rows(values_only=True)]
    wb.close()
    assert [r[0] for r in rows] == ["one.xlsx", "甲", "乙", "two.xlsx", "丙"]


def test_merge_ignores_lock_files_and_its_own_output(tmp_path: Path):
    folder = tmp_path / "表"
    make_book(folder / "one.xlsx", {"S1": [["甲"]]})
    make_book(folder / "~$one.xlsx", {"S1": [["锁文件"]]})
    make_book(folder / merge.DEFAULT_NAME, {"S1": [["上一次的结果"]]})

    merge.run(folder)
    from openpyxl import load_workbook

    wb = load_workbook(folder / merge.DEFAULT_NAME)
    rows = [list(r) for r in wb[merge.SHEET_NAME].iter_rows(values_only=True)]
    wb.close()
    assert rows == [["甲"]], "锁文件和输出文件自己都不该被合进来"


def test_xls_is_listed_but_skipped_with_a_reason(tmp_path: Path):
    folder = tmp_path / "表"
    make_book(folder / "one.xlsx", {"S1": [["甲"]]})
    (folder / "老格式.xls").write_bytes(b"\xd0\xcf\x11\xe0")

    report = merge.run(folder)
    skipped = [i for i in report.items if i.name == "老格式.xls"]
    assert skipped and skipped[0].status == SKIPPED
    assert "xls" in skipped[0].detail


def test_merge_plan_writes_nothing(tmp_path: Path):
    folder = tmp_path / "表"
    make_book(folder / "one.xlsx", {"S1": [["甲"]]})
    report = merge.plan(folder)
    assert report.counts[WOULD] == 1
    assert not (folder / merge.DEFAULT_NAME).exists()


def test_merge_on_a_folder_without_workbooks_says_so(tmp_path: Path):
    folder = tmp_path / "空"
    folder.mkdir()
    report = merge.plan(folder)
    assert report.items[0].status == SKIPPED
    assert "没有" in report.items[0].detail


# ------------------------------------------------------------------ 自检
def test_selfcheck_reports_the_ledger_and_pillow(sandbox):
    from temu_bot import toolbox

    info = toolbox.selfcheck(sandbox["cfg"])
    assert info["ledger"]["merged_total"] == 0
    assert info["pillow"]["avif"] is True
    assert info["pillow"]["webp"] is True
    assert set(info["modes"]) == set(rename.MODES)
    assert set(info["targets"]) == {"png", "jpg", "webp", "avif"}
